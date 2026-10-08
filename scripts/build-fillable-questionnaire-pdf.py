# Builds the fillable, locked PDF version of the discovery questionnaire from its .docx.
# Usage: pip install reportlab pypdf python-docx
#   python3 scripts/build-fillable-questionnaire-pdf.py <questionnaire.docx> <output.pdf> docs/pre-implementation/afronovation-logo.png
# Prints a newly generated owner password (needed to edit the PDF itself); keep it out of the repo.
"""Convert the discovery questionnaire .docx into a fillable, locked PDF."""
import sys, re, html, secrets
import docx
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Table,
                                TableStyle, Flowable, PageBreak, Image, KeepTogether, CondPageBreak)
from pypdf import PdfReader, PdfWriter
from pypdf.constants import UserAccessPermissions as P

SRC, OUT, LOGO = sys.argv[1], sys.argv[2], sys.argv[3]
GREEN = colors.HexColor('#14532D'); GOLD = colors.HexColor('#B8860B'); GREY = colors.HexColor('#5F6B66')
LIGHT = colors.HexColor('#EEF3EF'); LINE = colors.HexColor('#C9D3CD'); FIELD = colors.HexColor('#F4F8FC')
FIELD_BORDER = colors.HexColor('#9DB4CC')

body = ParagraphStyle('b', fontName='Helvetica', fontSize=9.6, leading=13, spaceAfter=6)
intro = ParagraphStyle('i', parent=body, fontName='Helvetica-Oblique', textColor=GREY)
h1 = ParagraphStyle('h1', fontName='Helvetica-Bold', fontSize=13.5, leading=17, textColor=GREEN, spaceBefore=14, spaceAfter=6)
cellS = ParagraphStyle('c', parent=body, fontSize=9, leading=11.5, spaceAfter=0)
headS = ParagraphStyle('hd', parent=cellS, fontName='Helvetica-Bold', textColor=colors.white)
bulletS = ParagraphStyle('bl', parent=body, leftIndent=16, bulletIndent=4, spaceAfter=3)
alignS = ParagraphStyle('al', parent=body, fontSize=9.2, leading=12.5, backColor=LIGHT, borderPadding=(5, 6, 5, 8),
                        leftIndent=4, spaceBefore=4, spaceAfter=10)

def markup(par):
    out = []
    for r in par.runs:
        s = html.escape(r.text)
        if not s: continue
        if r.bold: s = f'<b>{s}</b>'
        if r.italic: s = f'<i>{s}</i>'
        out.append(s)
    return ''.join(out)

class Field(Flowable):
    """A fillable multiline text field sized to its table cell."""
    def __init__(self, name, width, height, tip):
        super().__init__(); self.name, self.w, self.h, self.tip = name, width, height, tip
    def wrap(self, aw, ah): return self.w, self.h
    def draw(self):
        x, y = self.canv.absolutePosition(0, 0)
        self.canv.acroForm.textfield(name=self.name, tooltip=self.tip, x=x, y=y, width=self.w, height=self.h,
            fieldFlags='multiline' if self.h > 20 else '', fontName='Helvetica', fontSize=9 if self.h > 20 else 0,
            borderStyle='solid', borderWidth=0.6, borderColor=FIELD_BORDER, fillColor=FIELD, textColor=colors.black,
            forceBorder=True)

def tstyle(header=True):
    st = [('GRID', (0, 0), (-1, -1), 0.5, LINE), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
          ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
          ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]
    if header: st += [('BACKGROUND', (0, 0), (-1, 0), GREEN)]
    return TableStyle(st)

W = A4[0] - 2 * 20 * mm
d = docx.Document(SRC)
story, para_i, tbl_i = [], 0, 0
paras, tables = d.paragraphs, d.tables
nfields = 0

def widths(t):
    ws = [c.width for c in t.rows[0].cells]
    tot = sum(ws); return [W * w / tot for w in ws]

def render_table(t):
    global nfields
    rows = [[c.text for c in r.cells] for r in t.rows]
    cw = widths(t); hdr = rows[0]
    if hdr[:3] == ['No.', 'Question', 'Response']:
        data = [[Paragraph(h, headS) for h in hdr]]
        for no, q, _ in rows[1:]:
            qp = Paragraph(html.escape(q), cellS)
            qh = qp.wrap(cw[1] - 10, 1000)[1]
            data.append([Paragraph(f'<b>{no}</b>', cellS), qp, Field(no, cw[2] - 10, max(qh, 46), f'{no}: {q[:200]}')])
            nfields += 1
        tb = Table(data, colWidths=cw, repeatRows=1); tb.setStyle(tstyle()); return tb
    if hdr[0] == 'Ministry / Institution':  # Annex A
        cw = [W * x / 9638 for x in [2050, 1600, 1450, 1450, 1888, 1200]]
        data = [[Paragraph(h, headS) for h in hdr]]
        for i, r in enumerate(rows[1:]):
            row = []
            for j, v in enumerate(r):
                if v: row.append(Paragraph(html.escape(v), cellS))
                else:
                    row.append(Field(f'AnnexA.r{i+1}.c{j+1}', cw[j] - 10, 26, f'Annex A row {i+1}: {hdr[j]}')); nfields += 1
            data.append(row)
        tb = Table(data, colWidths=cw, repeatRows=1); tb.setStyle(tstyle()); return tb
    if hdr[0] == 'Area of responsibility':  # Annex B
        data = [[Paragraph(h, headS) for h in hdr]]
        for i, r in enumerate(rows[1:]):
            row = [Paragraph(html.escape(r[0]), cellS)]
            for j in range(1, len(r)):
                row.append(Field(f'AnnexB.{r[0]}.{hdr[j]}', cw[j] - 10, 26, f'{r[0]}: {hdr[j]}')); nfields += 1
            data.append(row)
        tb = Table(data, colWidths=cw); tb.setStyle(tstyle()); return tb
    if len(hdr) == 2 and hdr[0] == 'To':  # cover metadata
        data = []
        for k, v in rows:
            if k == 'Responses requested by':
                data.append([Paragraph(f'<b>{k}</b>', cellS), Field('ResponsesRequestedBy', 60 * mm, 15, 'Response date')]); nfields += 1
            else:
                data.append([Paragraph(f'<b>{k}</b>', cellS), Paragraph(html.escape(v), cellS)])
        tb = Table(data, colWidths=cw); st = tstyle(False); st.add('BACKGROUND', (0, 0), (0, -1), LIGHT)
        st.add('TEXTCOLOR', (0, 0), (0, -1), GREEN); tb.setStyle(st); return tb
    data = [[Paragraph(h, headS) for h in hdr]] + [[Paragraph(html.escape(v), cellS) for v in r] for r in rows[1:]]
    tb = Table(data, colWidths=cw, repeatRows=1); tb.setStyle(tstyle()); return tb

# Walk the body in document order
pmap = {p._p: p for p in paras}; tmap = {t._tbl: t for t in tables}
first = True; numbered = 0; cover_done = False
for el in d.element.body.iterchildren():
    if el.tag == qn('w:tbl'):
        story.append(render_table(tmap[el])); story.append(Spacer(1, 6)); continue
    if el.tag != qn('w:p'): continue
    p = pmap[el]; txt = p.text.strip()
    if el.find('.//' + qn('w:br')) is not None and not txt:
        story.append(PageBreak()); continue
    if not txt: story.append(Spacer(1, 4)); continue
    if txt == 'AFRONOVATION':
        story += [Spacer(1, 30), Image(LOGO, width=62 * mm, height=12.4 * mm, hAlign='LEFT'), Spacer(1, 10)]; continue
    if txt == 'ZidaProject' and not cover_done:
        story.append(Paragraph('ZidaProject', ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=26, leading=32, textColor=GREEN))); continue
    if txt == 'Pre-Implementation Discovery Questionnaire' and not cover_done:
        story.append(Paragraph(txt, ParagraphStyle('st', fontName='Helvetica', fontSize=17, leading=22, textColor=GREEN, spaceAfter=6))); continue
    if txt.startswith('Zimbabwe Digital Investment') and not cover_done:
        story.append(Paragraph(txt, ParagraphStyle('s2', parent=body, textColor=GREY, fontSize=10.5, spaceAfter=4,
                     borderPadding=(0, 0, 6, 0))))
        story.append(Table([['']], colWidths=[W], rowHeights=[2], style=[('LINEBELOW', (0, 0), (-1, -1), 1.5, GOLD)]))
        story.append(Spacer(1, 12)); cover_done = True; continue
    style = p.style.name
    if style.startswith('Heading'):
        story.append(CondPageBreak(60 * mm)); story.append(Paragraph(html.escape(txt), h1))
        story.append(Table([['']], colWidths=[W], rowHeights=[1], style=[('LINEABOVE', (0, 0), (-1, -1), 0.8, GOLD)]))
        if txt == 'How to Respond':
            story.append(Paragraph('Type your responses directly into the shaded fields, then save the file. Adobe Acrobat Reader (free) is recommended. The document text is protected; only the response fields can be edited.', bulletS, bulletText='•'))
        numbered = 0; continue
    numPr = el.find('.//' + qn('w:numPr'))
    if numPr is not None:
        numId = numPr.find(qn('w:numId')).get(qn('w:val'))
        fmt = d.part.numbering_part.numbering_definitions._numbering.xpath(
            f'.//w:num[@w:numId="{numId}"]')[0]
        absId = fmt.find(qn('w:abstractNumId')).get(qn('w:val'))
        lvl = d.part.numbering_part.numbering_definitions._numbering.xpath(
            f'.//w:abstractNum[@w:abstractNumId="{absId}"]/w:lvl/w:numFmt')[0].get(qn('w:val'))
        if lvl == 'decimal':
            numbered += 1; story.append(Paragraph(markup(p), bulletS, bulletText=f'{numbered}.'))
        else:
            story.append(Paragraph(markup(p), bulletS, bulletText='•'))
        continue
    if txt.startswith('NDAP alignment'):
        story.append(Paragraph(markup(p).replace('<b>NDAP alignment', '<font color="#14532D"><b>NDAP alignment').replace(' — </b>', ' — </b></font>', 1), alignS)); continue
    if all(r.italic for r in p.runs if r.text.strip()):
        story.append(Paragraph(html.escape(txt), intro)); continue
    story.append(Paragraph(markup(p), body))

def deco(c, doc):
    c.saveState(); c.setFont('Helvetica', 7.5); c.setFillColor(GREY)
    if doc.page > 1:
        c.drawRightString(A4[0] - 20 * mm, A4[1] - 12 * mm, 'ZidaProject  |  Pre-Implementation Discovery Questionnaire')
    c.drawCentredString(A4[0] / 2, 10 * mm, f'Confidential  |  Afronovation  |  Page {doc.page}')
    c.restoreState()

tmp = OUT + '.tmp.pdf'
doc = BaseDocTemplate(tmp, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=17 * mm,
                      title='ZidaProject Pre-Implementation Discovery Questionnaire', author='Afronovation, Inc.',
                      subject='Pre-Implementation Discovery Questionnaire for ZIDA')
doc.addPageTemplates([PageTemplate('p', [Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id='f')], onPage=deco)])
doc.build(story)

# Lock: allow filling, printing and copying text; block content editing.
owner = secrets.token_urlsafe(18)
w = PdfWriter(clone_from=tmp)
w.encrypt(user_password='', owner_password=owner, algorithm='AES-256',
          permissions_flag=P.PRINT | P.FILL_FORM_FIELDS | P.EXTRACT_TEXT_AND_GRAPHICS | P.PRINT_TO_REPRESENTATION | P.EXTRACT)
with open(OUT, 'wb') as f: w.write(f)
import os; os.remove(tmp)
print('fields:', nfields, 'pages:', len(PdfReader(OUT).pages)); print('OWNER_PASSWORD', owner)
